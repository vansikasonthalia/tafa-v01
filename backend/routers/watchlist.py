"""
Watchlist router — upload, list, and delete tickers.
Tickers persist permanently until explicitly removed.
"""

import io
import csv
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import func

from ..database import get_db
from ..models import Watchlist

router = APIRouter(prefix="/api/watchlist", tags=["watchlist"])


@router.post("/upload")
async def upload_watchlist(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    Upload a .csv or .xlsx file containing stock tickers.
    Expects columns: 'symbol' (or 'ticker'), optionally 'company name' and 'industry'.
    New tickers are added; duplicates are silently skipped.
    """
    filename = file.filename or ""
    content = await file.read()

    # Each entry: { "ticker": str, "company_name": str|None, "industry": str|None }
    entries = []

    if filename.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
            ws = wb.active

            headers = [str(cell.value).strip().lower() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

            ticker_col = None
            if "ticker" in headers:
                ticker_col = headers.index("ticker")
            elif "symbol" in headers:
                ticker_col = headers.index("symbol")
            else:
                ticker_col = 0

            name_col = headers.index("company name") if "company name" in headers else None
            industry_col = headers.index("industry") if "industry" in headers else None

            for row in ws.iter_rows(min_row=2):
                val = row[ticker_col].value
                if val and str(val).strip():
                    entries.append({
                        "ticker": str(val).strip().upper(),
                        "company_name": str(row[name_col].value).strip() if name_col is not None and row[name_col].value else None,
                        "industry": str(row[industry_col].value).strip() if industry_col is not None and row[industry_col].value else None,
                    })

            wb.close()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {str(e)}")

    elif filename.endswith(".csv") or True:  # Default to CSV
        try:
            text = content.decode("utf-8")
            reader = csv.reader(io.StringIO(text))

            headers = [h.strip().lower() for h in next(reader, [])]

            ticker_col = None
            if "ticker" in headers:
                ticker_col = headers.index("ticker")
            elif "symbol" in headers:
                ticker_col = headers.index("symbol")
            else:
                ticker_col = 0

            name_col = headers.index("company name") if "company name" in headers else None
            industry_col = headers.index("industry") if "industry" in headers else None

            for row in reader:
                if row and len(row) > ticker_col:
                    val = row[ticker_col].strip()
                    if val and val.upper() not in ("TICKER", "SYMBOL", ""):
                        company_name = None
                        industry = None
                        if name_col is not None and len(row) > name_col:
                            company_name = row[name_col].strip() or None
                        if industry_col is not None and len(row) > industry_col:
                            industry = row[industry_col].strip() or None

                        entries.append({
                            "ticker": val.upper(),
                            "company_name": company_name,
                            "industry": industry,
                        })

        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse CSV file: {str(e)}")

    if not entries:
        raise HTTPException(status_code=400, detail="No tickers found in the uploaded file.")

    # Auto-append .NS for Indian stocks if no suffix exists (Yahoo Finance compatibility)
    for entry in entries:
        if "." not in entry["ticker"]:
            entry["ticker"] = f"{entry['ticker']}.NS"

    # Add new tickers, skip duplicates
    added = 0
    skipped = 0
    for entry in entries:
        existing = db.query(Watchlist).filter(func.upper(Watchlist.ticker) == entry["ticker"]).first()
        if existing:
            skipped += 1
        else:
            db.add(Watchlist(
                ticker=entry["ticker"],
                company_name=entry["company_name"],
                industry=entry["industry"],
            ))
            added += 1

    db.commit()

    total = db.query(Watchlist).count()

    return {
        "added": added,
        "skipped": skipped,
        "total": total,
    }


@router.get("")
def get_watchlist(db: Session = Depends(get_db)):
    """Return all tickers in the watchlist."""
    items = db.query(Watchlist).order_by(Watchlist.created_at.desc()).all()
    return {
        "total": len(items),
        "tickers": [item.to_dict() for item in items],
    }


@router.delete("")
def clear_watchlist(db: Session = Depends(get_db)):
    """Remove ALL tickers from the watchlist."""
    count = db.query(Watchlist).delete()
    db.commit()
    return {"deleted": count}


@router.delete("/{ticker}")
def remove_ticker(ticker: str, db: Session = Depends(get_db)):
    """Remove a single ticker from the watchlist."""
    item = db.query(Watchlist).filter(func.upper(Watchlist.ticker) == ticker.upper()).first()
    if not item:
        raise HTTPException(status_code=404, detail=f"Ticker '{ticker}' not found in watchlist.")
    db.delete(item)
    db.commit()
    return {"deleted": ticker.upper()}
